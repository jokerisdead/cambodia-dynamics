#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Cambodia weekly dynamics Word report from JSON data."""

import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Load data
with open(r'D:\龙虾\automation-claw-2026-05-14-task-1\data.json',
          'r', encoding='utf-8') as f:
    items = json.load(f)

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

# Normal style font
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# Title
title = doc.add_heading('', level=0)
run = title.add_run('柬埔寨2026年第21周动态')
run.font.name = '黑体'
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0, 0, 0)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Subtitle
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('（2026年5月17日—5月23日）')
run.font.name = '楷体'
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(80, 80, 80)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')

doc.add_paragraph('─' * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER

# Editor's note
note = doc.add_paragraph()
run = note.add_run('【编者按】')
run.bold = True
run.font.name = '黑体'
run.font.size = Pt(12)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
run = note.add_run(
    '本期动态聚焦2026年第21周（5月17日至5月23日）柬埔寨国内热点事件，'
    '重点关注宗教领域动态及中柬经贸合作进展。'
    '本周联合国卫塞节在中国无锡举办，柬埔寨佛教界应邀参与；'
    '波贝僧侣涉电诈大案持续发酵引发宗教治理反思；'
    '中柬在农业贸易、新能源投资、旅游合作等领域延续强劲合作势头。'
)
run.font.name = '宋体'
run.font.size = Pt(12)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.add_paragraph('')

# Build each item
for item in items:
    # Sub-heading
    h = doc.add_heading('', level=2)
    run = h.add_run(f'{item["num"]}. {item["title"]}')
    run.font.name = '黑体'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    # Date & source
    p = doc.add_paragraph()
    run = p.add_run(f'时间：{item["date"]}  |  来源：{item["source"]}')
    run.font.name = '楷体'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.italic = True
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')

    # Content
    p = doc.add_paragraph(item['content'])
    for run in p.runs:
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(22)

    # Comment
    p = doc.add_paragraph()
    run = p.add_run('【简评】')
    run.bold = True
    run.font.name = '黑体'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(153, 51, 0)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run = p.add_run(item['comment'])
    run.font.name = '楷体'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(80, 80, 80)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(22)

    doc.add_paragraph('')

# Footer
doc.add_paragraph('─' * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('九江学院柬埔寨研究中心  编制')
run.font.name = '楷体'
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(100, 100, 100)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')

footer2 = doc.add_paragraph()
footer2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer2.add_run('2026年5月23日')
run.font.name = '楷体'
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(100, 100, 100)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')

# Save
output_path = r'D:\龙虾\automation-claw-2026-05-14-task-1\柬埔寨2026年第21周动态.docx'
doc.save(output_path)
print(f'文档已保存至：{output_path}')
